import io
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from PIL import Image

from presentation_maker_offline.export import export_project_pptx
from presentation_maker_offline.project_document import new_project_document
from presentation_maker_offline.source_import import parse_outline_text
from presentation_maker_offline.layout_design import apply_slide_layout


def test_export_project_creates_editable_text_and_keeps_coordinates(tmp_path):
    document = new_project_document("季度報告")
    document["slides"] = [{
        "id": "slide-1", "order": 0, "title": "營運摘要",
        "elements": [{"id": "text-1", "type": "text", "text": "營收成長 12%", "x": .2, "y": .3, "width": .5, "height": .2}],
    }]
    output = tmp_path / "editable.pptx"
    export_project_pptx(output, document)

    imported = Presentation(output)
    assert len(imported.slides) == 1
    texts = [shape.text for shape in imported.slides[0].shapes if getattr(shape, "has_text_frame", False)]
    assert "營運摘要" in texts
    assert "營收成長 12%" in texts
    body = next(shape for shape in imported.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text == "營收成長 12%")
    assert .19 < body.left / imported.slide_width < .21
    assert .49 < body.width / imported.slide_width < .51


def test_export_project_embeds_project_images_and_rejects_path_escape(tmp_path):
    asset_root = tmp_path / "assets"
    asset_root.mkdir()
    Image.new("RGB", (40, 20), "tomato").save(asset_root / "cover.png")
    document = new_project_document("封面")
    document["slides"] = [{"id": "slide-1", "order": 0, "title": "封面", "elements": [
        {"id": "image-1", "type": "image", "asset_path": "cover.png", "x": .1, "y": .2, "width": .4, "height": .4},
    ]}]
    output = tmp_path / "with-image.pptx"
    export_project_pptx(output, document, asset_root=asset_root)
    imported = Presentation(output)
    pictures = [shape for shape in imported.slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 1
    assert pictures[0].image.ext == "png"

    document["slides"][0]["elements"][0]["asset_path"] = "../outside.png"
    import pytest
    with pytest.raises(ValueError, match="inside the project asset folder"):
        export_project_pptx(tmp_path / "unsafe.pptx", document, asset_root=asset_root)


def test_user_outline_exports_all_pages_with_selected_theme_and_readable_text(tmp_path):
    chunks = ["**認識大腸癌**"]
    for page in range(1, 22):
        chunks.append(f"# 第{page}頁｜第{page}頁主題\n- 早期發現與健康管理重點\n- 定期篩檢及日常照護")
    title, slides, source = parse_outline_text("\n\n".join(chunks))
    document = new_project_document(title)
    document["slides"] = slides
    document["sources"] = [source]
    for slide in slides:
        slide["source_id"] = source["id"]

    output = export_project_pptx(tmp_path / "colon-health.pptx", document, "自然療癒")
    reopened = Presentation(output)
    assert len(reopened.slides) == 21
    first = reopened.slides[0]
    assert first.background.fill.fore_color.rgb == RGBColor(0xF2, 0xF8, 0xF1)
    first_text = "\n".join(shape.text for shape in first.shapes if shape.has_text_frame)
    assert "第1頁主題" in first_text
    assert "定期篩檢及日常照護" in first_text
    last_text = "\n".join(shape.text for shape in reopened.slides[-1].shapes if shape.has_text_frame)
    assert "第21頁主題" in last_text
    assert "21 / 21" in last_text


def test_two_export_modes_preserve_selected_layout_and_image(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from presentation_maker_offline.raster_export import export_project_image_pptx

    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (120, 80), "tomato").save(assets / "illustration.png")
    document = new_project_document("健康簡報")
    document["slides"] = [{
        "id": "s1", "order": 0, "title": "健康簡報",
        "elements": [
            {"id": "t1", "type": "text", "text": "生活與健康", "x": .08, "y": .24, "width": .84, "height": .58},
            {"id": "i1", "type": "image", "asset_path": "illustration.png", "x": .58, "y": .22, "width": .36, "height": .60},
        ],
    }]
    apply_slide_layout(document["slides"][0], "左圖右文", cover=True)
    editable = Presentation(export_project_pptx(tmp_path / "editable-layout.pptx", document, asset_root=assets))
    image_only = Presentation(export_project_image_pptx(tmp_path / "image-layout.pptx", document, asset_root=assets))

    editable_shapes = editable.slides[0].shapes
    editable_body = next(shape for shape in editable_shapes if shape.has_text_frame and shape.text == "生活與健康")
    editable_picture = next(shape for shape in editable_shapes if shape.shape_type == 13)
    assert editable_picture.left < editable_body.left
    assert any(shape.has_text_frame and shape.text == "健康簡報" for shape in editable_shapes)

    raster_shapes = list(image_only.slides[0].shapes)
    assert len(raster_shapes) == 1
    assert raster_shapes[0].shape_type == 13
    with Image.open(io.BytesIO(raster_shapes[0].image.blob)) as rendered:
        assert rendered.size == (1920, 1080)
        assert rendered.getpixel((480, 520))[:3] == (255, 99, 71)
        assert rendered.getpixel((1800, 950))[:3] != (255, 99, 71)


def test_top_image_layout_uses_center_crop_in_editable_output(tmp_path):
    assets = tmp_path / "assets"
    assets.mkdir()
    art = Image.new("RGB", (100, 100), "green")
    art.save(assets / "square.png")
    document = new_project_document("排版")
    slide = {"id": "s1", "order": 0, "title": "排版", "elements": [
        {"id": "i1", "type": "image", "asset_path": "square.png"},
        {"id": "t1", "type": "text", "text": "重要內容"},
    ]}
    document["slides"] = [slide]
    apply_slide_layout(slide, "上圖下文", cover=True)
    deck = Presentation(export_project_pptx(tmp_path / "top-image.pptx", document, asset_root=assets))
    picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == 13)
    assert picture.width / picture.height > 4
    with Image.open(io.BytesIO(picture.image.blob)) as embedded:
        assert embedded.width / embedded.height > 4
