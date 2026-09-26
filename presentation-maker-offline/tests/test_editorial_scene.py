import pytest
import hashlib
from PIL import Image
from presentation_maker_offline.editorial_scene import eligibility_scene
from presentation_maker_offline.editorial_scene import checked_visual
from presentation_maker_offline.editorial_scene import attach_scene, current_scene


def test_scene_dual_export_and_preview(tmp_path):
    import os
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from pptx import Presentation
    from PySide6.QtWidgets import QApplication
    from presentation_maker_offline.project_document import new_project_document
    from presentation_maker_offline.export import export_project_pptx
    from presentation_maker_offline.raster_export import export_project_image_pptx
    from presentation_maker_offline.qt_ui import SlidePreview, slide_preview_payload
    app = QApplication.instance() or QApplication([])
    lines = ['45至74歲民眾', '40至44歲，父母、子女或兄弟姊妹曾罹患大腸癌者', '每2年補助1次糞便潛血檢查']
    scene = eligibility_scene('誰符合公費大腸癌篩檢？', lines)
    doc = new_project_document('驗收')
    doc['slides'] = [dict(id='p17', title='誰符合公費大腸癌篩檢？', elements=[])]
    attach_scene(doc['slides'][0], scene)
    editable = export_project_pptx(tmp_path/'editable.pptx', doc)
    image = export_project_image_pptx(tmp_path/'image.pptx', doc, asset_root=tmp_path)
    native = Presentation(editable).slides[0]
    assert ''.join(s.text for s in native.shapes if s.has_text_frame) == ''.join(t['text'] for t in scene['texts'])
    assert all(s.has_text_frame for s in native.shapes)
    raster = Presentation(image).slides[0]
    assert len(raster.shapes) == 1
    assert raster.shapes[0].image.size == (1920, 1080)
    canvas = SlidePreview()
    canvas.set_slide(**slide_preview_payload(doc, 0, tmp_path))
    canvas.save_slide_image(tmp_path/'preview.png')
    assert (tmp_path/'preview.png').exists()
    canvas.close()
    doc['slides'][0]['title'] = '已改標題'
    with pytest.raises(ValueError, match='內容已變更'):
        current_scene(doc['slides'][0])


def test_visual_requires_review_bound_to_image(tmp_path):
    path = tmp_path / 'visual.png'
    Image.new('RGB', (32, 32)).save(path)
    item = {'path': str(path)}
    with pytest.raises(ValueError, match='accepted'):
        checked_visual(item)
    item['review'] = dict(verdict='accepted', image_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    assert checked_visual(item).width() == 32
    Image.new('RGB', (64, 32)).save(path)
    with pytest.raises(ValueError, match='changed'):
        checked_visual(item)


def test_exact_copy_survives_typographic_hierarchy():
    lines = ['45至74歲民眾', '40至44歲，父母、子女或兄弟姊妹曾罹患大腸癌者', '每2年補助1次糞便潛血檢查']
    scene = eligibility_scene('誰符合公費大腸癌篩檢？', lines)
    assert ''.join(item['text'] for item in scene['texts'][1:]) == ''.join(lines)
    with pytest.raises(ValueError):
        eligibility_scene('title', ['one'])


def test_comparison_design_is_used_and_regenerated_after_edit():
    from presentation_maker_offline.layout_design import auto_design_slide, apply_slide_layout
    slide = dict(title='比較', elements=[dict(id='body', type='text',
                 text='迷思：甲\n正確觀念：乙\n迷思：丙\n正確觀念：丁')])
    auto_design_slide(slide, 1)
    assert current_scene(slide)['generated_by'] == 'semantic-comparison'
    slide['elements'][0]['text'] = slide['elements'][0]['text'].replace('乙', '更新')
    auto_design_slide(slide, 1)
    assert '更新' in ''.join(t['text'] for t in current_scene(slide)['texts'])
    apply_slide_layout(slide, '左圖右文', cover=False)
    assert 'editorial_scene' not in slide
