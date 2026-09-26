from pptx import Presentation

from presentation_maker_offline.export import export_project_pptx
from presentation_maker_offline.layout_design import auto_design_slide, render_text_blocks
from presentation_maker_offline.project_document import new_project_document
from presentation_maker_offline.source_import import parse_outline_text


def test_content_driven_layout_for_colorectal_health_outline(tmp_path):
    outline = """# 認識大腸癌
## 第1頁｜封面：認識大腸癌
- 早期發現、早期治療
## 第2頁｜為什麼需要認識大腸癌？
- 早期可能沒有明顯症狀
- 定期篩檢有助於提早發現
## 第3頁｜大腸癌不是突然發生
- 正常黏膜 → 瘜肉 → 癌前病變 → 大腸癌
- 這個過程需要時間
## 第4頁｜無法改變的危險因子
- 年齡增加
- 大腸癌家族史
- 曾有大腸瘜肉
- 遺傳性疾病
- 發炎性腸道疾病
- 有危險因子不代表一定會罹癌
"""
    _title, slides, source = parse_outline_text(outline)
    kinds = [auto_design_slide(slide, index) for index, slide in enumerate(slides)]
    assert kinds == ["封面主視覺", "重點敘述", "流程步驟", "雙欄重點"]
    assert all(slide["auto_layout"] for slide in slides)
    flow_blocks = render_text_blocks(slides[2])
    assert len(flow_blocks) == 5
    assert all(block.get("card") for block in flow_blocks[:4])
    double_blocks = render_text_blocks(slides[3])
    assert len(double_blocks) == 2
    original_lines = slides[3]["elements"][0]["text"].splitlines()
    assert "\n".join(block["text"] for block in double_blocks).splitlines() == original_lines

    document = new_project_document("認識大腸癌")
    document["slides"] = slides
    document["sources"] = [source]
    output = tmp_path / "auto-design.pptx"
    export_project_pptx(output, document)
    pptx = Presentation(output)
    assert len(pptx.slides) == 4
    exported_flow = "\n".join(shape.text for shape in pptx.slides[2].shapes if shape.has_text_frame)
    assert all(term in exported_flow for term in ("正常黏膜", "瘜肉", "癌前病變", "大腸癌", "這個過程需要時間"))


def test_manual_override_does_not_get_replaced_during_render():
    slide = {"id": "slide", "title": "六項重點", "auto_layout": False,
             "elements": [{"id": "text", "type": "text", "text": "\n".join(str(i) for i in range(6)),
                           "x": .08, "y": .24, "width": .84, "height": .58}]}
    assert len(render_text_blocks(slide)) == 1


def test_myths_are_arranged_as_comparison_without_losing_facts():
    slide = {"id": "myths", "title": "破解常見迷思", "elements": [
        {"id": "body", "type": "text", "text":
         "迷思：沒有症狀就不用檢查\n正確觀念：早期可能沒有症狀\n"
         "迷思：血便一定是痔瘡\n正確觀念：血便應由醫師評估"}]}
    assert auto_design_slide(slide, 19) == "迷思對照"
    blocks = render_text_blocks(slide)
    assert len(blocks) == 2
    assert "沒有症狀就不用檢查" in blocks[0]["text"]
    assert "血便應由醫師評估" in blocks[1]["text"]
