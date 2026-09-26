import os
import time
from pathlib import Path

import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pptx import Presentation
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest

from presentation_maker_offline.qt_ui import PresentationStudio


@pytest.fixture
def studio(tmp_path):
    app = QApplication.instance() or QApplication([])
    window = PresentationStudio(tmp_path)
    yield app, window
    window.close()


def test_qt_outline_clipboard_preview_edit_save_reopen_and_export(studio, tmp_path, monkeypatch):
    app, window = studio
    app.clipboard().setText("# 衛教簡報\n## 認識疾病\n- 早期發現\n## 預防行動\n- 定期篩檢")
    window._start_outline()
    window._paste_clipboard()
    app.processEvents()

    assert window.slide_preview.count() == 2
    assert window.continue_button.isEnabled()
    assert "2 頁" in window.parse_status.text()

    window._continue_settings()
    window.settings_title.setText("大腸癌衛教")
    window.style_choice.setCurrentText("自然療癒")
    window._create_project()
    assert window.stack.currentWidget() is window.editor_page
    assert window.slide_list.count() == 2
    assert window.slide_canvas.scene.items()

    window.slide_list.setCurrentRow(1)
    window.slide_text.setPlainText("• 每兩年接受篩檢\n• 有警訊應就醫")
    window._apply_slide_text()
    assert window.revision == 2
    reopened = window.store.load_document(window.project_id)[1]
    assert "每兩年接受篩檢" in reopened["slides"][1]["elements"][0]["text"]

    output = tmp_path / "health.pptx"
    from presentation_maker_offline import qt_ui
    monkeypatch.setattr(qt_ui.QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(output), "PowerPoint (*.pptx)"))
    monkeypatch.setattr(qt_ui.QMessageBox, "information", lambda *_args, **_kwargs: None)
    window.export_project()
    assert len(Presentation(output).slides) == 2
    assert window.store.load_document(window.project_id)[1]["settings"]["style"] == "自然療癒"


def test_native_outline_paste_supports_macos_keyboard_and_context_menu(studio):
    app, window = studio
    window._start_outline()
    app.clipboard().setText("# 測試簡報\n## 第一頁\n- 可貼上")
    window.outline_input.setFocus()
    QTest.keyClick(window.outline_input, Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)
    app.processEvents()

    assert "可貼上" in window.outline_input.toPlainText()
    menu = window.outline_input.createStandardContextMenu()
    try:
        labels = [action.text().replace("&", "").lower() for action in menu.actions()]
        assert any("貼上" in label or "paste" in label for label in labels)
    finally:
        menu.deleteLater()


def test_qt_invalid_outline_keeps_next_step_disabled(studio):
    app, window = studio
    window._start_outline()
    window.outline_input.setPlainText("   ")
    app.processEvents()
    assert window.slide_preview.count() == 0
    assert not window.continue_button.isEnabled()
    assert "請先貼上" in window.parse_status.text()


def test_long_real_world_outline_with_table_and_joined_heading_remains_21_pages(studio):
    app, window = studio
    pages = [f"# 認識大腸癌｜21頁簡報大綱\n\n## 第1頁｜封面：認識大腸癌\n- 早期發現、早期治療"]
    pages.extend(f"## 第{index}頁｜衛教主題{index}\n- 重要重點{index}" for index in range(2, 7))
    pages.append(
        "## 第7頁｜瘜肉不等於大腸癌\n- 多數瘜肉屬良性\n- 部分可能癌變"
        "第8頁｜無法改變的危險因子\n- 年齡增加\n- 大腸癌家族史"
    )
    pages.extend(f"## 第{index}頁｜衛教主題{index}\n- 重要重點{index}" for index in range(9, 20))
    pages.append(
        "## 第20頁｜破解常見迷思\n"
        "| **常見迷思正確觀念** |                    |\n"
        "| ------------ | ------------------ |\n"
        "| 沒有症狀就不用檢查 | 早期可能沒有症狀 |\n"
        "## 第21頁｜守護腸道健康\n- 定期篩檢"
    )
    window._start_outline()
    window.outline_input.setPlainText("\n\n".join(pages) + "\n- \\")
    app.processEvents()

    assert window.slide_preview.count() == 21
    assert window.continue_button.isEnabled()
    assert window.parsed["slides"][7]["title"] == "無法改變的危險因子"
    assert "迷思：沒有症狀就不用檢查" in window.parsed["slides"][19]["elements"][0]["text"]
    assert "正確觀念：早期可能沒有症狀" in window.parsed["slides"][19]["elements"][0]["text"]


def test_ai_candidate_requires_explicit_accept_and_keeps_original_source(studio, monkeypatch):
    app, window = studio
    from presentation_maker_offline import qt_ui
    from presentation_maker_offline.source_import import parse_outline_text

    original = "# 原始簡報\n## 原始頁\n- 使用者提供的事實"
    candidate_text = "# AI 候選\n## 整理後的頁面\n- 保留原始事實"
    window._start_outline()
    window.outline_input.setPlainText(original)
    original_source = window.parsed["source"]
    _, candidate_slides, candidate_source = parse_outline_text(candidate_text)
    candidate_source.update({"origin": "model_generated_outline", "runtime": "mlx-serve 26.9.2"})
    monkeypatch.setattr(qt_ui.QDialog, "exec", lambda _self: qt_ui.QDialog.DialogCode.Accepted)

    window._review_ai_candidate({"title": "AI 候選", "slides": candidate_slides, "source": candidate_source, "runtime": "mlx-serve 26.9.2"})
    app.processEvents()

    assert window.original_outline_source["id"] == original_source["id"]
    assert window.model_outline_source["origin"] == "model_generated_outline"
    assert window.parsed["source"]["origin"] == "accepted_model_outline"
    assert window.parsed["source"]["model_source_id"] == window.model_outline_source["id"]
    assert "整理後的頁面" in window.outline_input.toPlainText()


def test_rejected_ai_candidate_does_not_replace_pasted_outline(studio, monkeypatch):
    _app, window = studio
    from presentation_maker_offline import qt_ui
    from presentation_maker_offline.source_import import parse_outline_text

    original = "# 原始簡報\n## 原始頁\n- 原始內容"
    candidate_text = "# 候選簡報\n## 候選頁\n- 新內容"
    window._start_outline()
    window.outline_input.setPlainText(original)
    _, slides, source = parse_outline_text(candidate_text)
    monkeypatch.setattr(qt_ui.QDialog, "exec", lambda _self: qt_ui.QDialog.DialogCode.Rejected)

    window._review_ai_candidate({"title": "候選簡報", "slides": slides, "source": source, "runtime": "mlx-serve 26.9.2"})

    assert window.outline_input.toPlainText() == original
    assert window.model_outline_source is None
    assert window.parsed["source"]["text"] == original


def _create_one_slide_project(app, window):
    window._start_outline()
    window.outline_input.setPlainText("# 測試簡報\n## 健康生活\n- 均衡飲食\n- 規律運動")
    window._continue_settings()
    window._create_project()
    app.processEvents()


def test_area_marking_uses_canvas_coordinates_and_survives_reopen(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    window.resize(1180, 780)
    window.show()
    app.processEvents()
    window.start_area_marking()
    app.processEvents()
    start = window.slide_canvas.mapFromScene(160, 400)
    end = window.slide_canvas.mapFromScene(600, 600)
    viewport = window.slide_canvas.viewport()
    QTest.mousePress(viewport, Qt.MouseButton.LeftButton, pos=start)
    QTest.mouseMove(viewport, end, 100)
    QTest.mouseRelease(viewport, Qt.MouseButton.LeftButton, pos=end)
    app.processEvents()

    assert window._pending_annotation is not None
    window.annotation_comment.setText("只修改這個區域")
    window.save_annotation()
    loaded_revision, reopened = window.store.load_document(window.project_id)
    mark = reopened["annotations"][0]
    assert loaded_revision == window.revision
    assert mark["comment"] == "只修改這個區域"
    assert mark["slide_id"] == reopened["slides"][0]["id"]
    assert mark["element_id"] == reopened["slides"][0]["elements"][0]["id"]
    assert len(mark["rect"]) == 4

    window._return_home()
    window._open_recent(window.recent_list.item(0))
    assert window.annotation_list.count() == 1
    window.annotation_list.setCurrentRow(0)
    window.annotation_comment.setText("留言可再編輯")
    window.save_annotation()
    assert len(window.document["annotations"]) == 1
    assert window.document["annotations"][0]["comment"] == "留言可再編輯"


def test_image_button_generates_saves_previews_and_exports_bound_asset(studio, tmp_path, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    from PIL import Image
    from presentation_maker_offline import qt_ui

    class FakeImageBackend:
        def __init__(self, *_args, **_kwargs):
            pass

        def health(self):
            return {"ok": True, "errors": []}

        def generate(self, prompt, *, cancel_event=None, progress_callback=None):
            generated = tmp_path / "generated.png"
            Image.new("RGB", (96, 64), (24, 142, 86)).save(generated)
            progress_callback(1, 2)
            return {"image_path": str(generated), "device": "mps", "vae_device": "mps"}

    monkeypatch.setattr(qt_ui, "LocalQwenImageBackend", FakeImageBackend)
    monkeypatch.setattr(qt_ui, "qwen_image21_readiness", lambda _path: {"ready": True, "missing": [], "incomplete": []})
    monkeypatch.setattr(qt_ui.QMessageBox, "critical", lambda *_args, **_kwargs: None)
    window.image_prompt.setPlainText("無文字的健康生活插圖")
    window.generate_current_image_button.click()

    deadline = time.monotonic() + 15
    while window._image_worker is not None and time.monotonic() < deadline:
        QTest.qWait(20)
    app.processEvents()

    assert window._image_worker is None
    slide = window.document["slides"][0]
    image_element = next(element for element in slide["elements"] if element.get("type") == "image")
    asset = tmp_path / "projects" / window.project_id / "assets" / image_element["asset_path"]
    assert asset.is_file()
    assert len([item for item in window.slide_canvas.scene.items() if hasattr(item, "pixmap")]) == 1
    assert window.document["image_generation_ledger"][0]["output_sha256"]
    assert window.store.attempts(window.project_id)[-1]["phase"] == "image_generation"

    output = tmp_path / "with-image.pptx"
    from presentation_maker_offline.export import export_project_pptx
    export_project_pptx(output, window.document, "清爽藍", asset_root=tmp_path / "projects" / window.project_id / "assets")
    exported = Presentation(output)
    assert len(exported.slides[0].shapes) >= 4
    assert any(shape.shape_type == 13 for shape in exported.slides[0].shapes)
    picture = next(shape for shape in exported.slides[0].shapes if shape.shape_type == 13)
    assert .57 <= picture.left / exported.slide_width <= .59
    text_shape = next(shape for shape in exported.slides[0].shapes if shape.has_text_frame and "均衡飲食" in shape.text)
    assert text_shape.width / exported.slide_width < .52


def test_export_mode_is_saved_and_save_failure_stops_export(studio, tmp_path, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    from presentation_maker_offline import qt_ui

    destination = tmp_path / "image-mode.pptx"
    monkeypatch.setattr(qt_ui.QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(destination), "PowerPoint (*.pptx)"))
    monkeypatch.setattr(qt_ui.QMessageBox, "information", lambda *_args, **_kwargs: None)
    window.output_mode_choice.setCurrentText("圖像式 PPTX")
    window.export_project()
    app.processEvents()
    assert len(Presentation(destination).slides[0].shapes) == 1
    assert window.store.load_document(window.project_id)[1]["settings"]["output_format"] == "圖像式 PPTX"

    blocked_destination = tmp_path / "should-not-export.pptx"
    monkeypatch.setattr(qt_ui.QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(blocked_destination), "PowerPoint (*.pptx)"))
    monkeypatch.setattr(window, "save_project", lambda: False)
    window.export_project()
    assert not blocked_destination.exists()


def test_layout_design_and_per_slide_image_prompt_survive_reopen(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    window.layout_choice.setCurrentText("左圖右文")
    window.apply_layout_button.click()
    window.image_prompt.setPlainText("綠色植物與日常運動，無字插圖")
    assert window.save_project()
    project_id = window.project_id

    window._return_home()
    window._open_recent(window.recent_list.item(0))
    app.processEvents()
    assert window.project_id == project_id
    assert window.layout_choice.currentText() == "左圖右文"
    assert window.image_prompt.toPlainText() == "綠色植物與日常運動，無字插圖"
    assert window.document["slides"][0]["layout"] == "左圖右文"


def test_editor_columns_resize_and_large_preview_contains_current_slide(studio, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    from presentation_maker_offline import qt_ui

    captured = {}
    def inspect_dialog(dialog):
        captured["title"] = dialog.windowTitle()
        previews = dialog.findChildren(qt_ui.SlidePreview)
        captured["preview_text"] = [item.toPlainText() for item in previews[0].scene.items()
                                    if hasattr(item, "toPlainText")]
        return qt_ui.QDialog.DialogCode.Accepted

    monkeypatch.setattr(qt_ui.QDialog, "exec", inspect_dialog)
    window.resize(900, 620)
    window.show()
    app.processEvents()
    assert window.zoom_preview_button.isVisible()
    assert window.export_button.isVisible()
    window.zoom_preview_button.click()
    assert "第 1 頁" in captured["title"]
    assert any("健康生活" in text for text in captured["preview_text"])


def test_reference_files_are_visible_persisted_and_do_not_change_slides(studio, tmp_path, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    reference = tmp_path / "衛教參考.txt"
    reference.write_text("篩檢資訊\n請由醫師評估", encoding="utf-8")
    from presentation_maker_offline import qt_ui
    monkeypatch.setattr(qt_ui.QFileDialog, "getOpenFileNames", lambda *_a, **_kw: ([str(reference)], ""))
    original_slides = window.store.load_document(window.project_id)[1]["slides"]

    window._add_reference_file()
    assert window.source_list.count() == 1
    assert "請由醫師評估" in window.source_preview.toPlainText()
    window.source_role.setCurrentText("修改依據")
    window.source_scope.setText("頁 1")
    window._save_source_settings()
    saved = window.store.load_document(window.project_id)[1]
    assert saved["slides"] == original_slides
    assert saved["sources"][-1]["role"] == "修改依據"
    assert saved["sources"][-1]["scope"] == "頁 1"

    window._return_home()
    window._open_recent(window.recent_list.item(0))
    assert window.source_list.count() == 1
    assert window.source_role.currentText() == "修改依據"
    assert "請由醫師評估" in window.source_preview.toPlainText()


def test_source_content_requires_confirm_and_keeps_citation(studio, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    from presentation_maker_offline import qt_ui
    from presentation_maker_offline.source_library import add_text_source
    source = add_text_source("可供引用的內容", window.document["sources"], name="衛教指南")
    assert window._append_library_source(source)
    before = len(window.document["slides"][0]["elements"])
    monkeypatch.setattr(qt_ui.QDialog, "exec", lambda _dialog: qt_ui.QDialog.DialogCode.Rejected)
    window._insert_selected_source_fragment()
    assert len(window.document["slides"][0]["elements"]) == before
    monkeypatch.setattr(qt_ui.QDialog, "exec", lambda _dialog: qt_ui.QDialog.DialogCode.Accepted)
    window._insert_selected_source_fragment()
    saved = window.store.load_document(window.project_id)[1]
    assert len(saved["slides"][0]["elements"]) == before + 1
    assert saved["slides"][0]["elements"][-1]["source_ref"]["source_id"] == source["id"]


def test_new_outline_auto_designs_each_page_and_reflows_after_text_edit(studio):
    app, window = studio
    window._start_outline()
    window.outline_input.setPlainText(
        "# 衛教簡報\n## 封面\n- 早期發現\n"
        "## 第2頁｜風險因素\n- 年齡\n- 家族史\n- 吸菸\n- 飲酒\n- 缺乏運動\n- 肥胖\n")
    window._continue_settings()
    window._create_project()
    assert window.document["slides"][1]["content_layout"] == "雙欄重點"
    assert window.document["slides"][1]["auto_layout"]
    window.slide_list.setCurrentRow(1)
    app.processEvents()
    assert "雙欄重點" in window.auto_layout_status.text()
    window.slide_text.setPlainText("• 只留下單一重點")
    window._apply_slide_text()
    assert window.document["slides"][1]["content_layout"] == "重點敘述"
    assert window.store.load_document(window.project_id)[1]["slides"][1]["content_layout"] == "重點敘述"


def test_revision_restore_creates_new_revision_without_erasing_history(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    original_revision = window.revision
    assert window.save_project()
    assert window.revision == original_revision
    window.slide_heading.setText("更改後標題")
    window._apply_slide_text()
    changed_revision = window.revision
    assert changed_revision > original_revision
    assert window.restore_project_revision(original_revision)
    assert window.revision > changed_revision
    assert window.document["slides"][0]["title"] == "健康生活"
    assert window.store.load_revision(window.project_id, changed_revision)[1]["slides"][0]["title"] == "更改後標題"
    assert window.store.revision_parent(window.project_id, window.revision) == original_revision


def test_text_edit_candidate_ui_accepts_only_after_review_and_preserves_history(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    slide = window.document["slides"][0]
    element = next(item for item in slide["elements"] if item["type"] == "text")
    original = element["text"]
    window._receive_text_edit_candidate({
        "slide_id": slide["id"], "element_id": element["id"], "base_revision": window.revision,
        "instruction": "改得更易懂", "original_text": original,
        "proposed_text": "• 均衡飲食與規律運動", "runtime": "test-local",
    })
    assert element["text"] == original
    assert window.edit_candidate_list.count() == 1
    assert "均衡飲食與規律運動" in window.edit_comparison.toPlainText()
    window.accept_selected_text_edit()
    saved = window.store.load_document(window.project_id)[1]
    assert saved["slides"][0]["elements"][0]["text"] == "• 均衡飲食與規律運動"
    assert saved["edit_candidates"][0]["status"] == "accepted"
    assert window.store.load_revision(window.project_id, 1)[1]["slides"][0]["elements"][0]["text"] == original


def test_text_edit_candidate_ui_rejects_without_changing_slide(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    slide = window.document["slides"][0]
    element = slide["elements"][0]
    original = element["text"]
    window._receive_text_edit_candidate({
        "slide_id": slide["id"], "element_id": element["id"], "base_revision": window.revision,
        "instruction": "縮短", "original_text": original, "proposed_text": "簡短版", "runtime": "test-local",
    })
    window.reject_selected_text_edit()
    saved = window.store.load_document(window.project_id)[1]
    assert saved["slides"][0]["elements"][0]["text"] == original
    assert saved["edit_candidates"][0]["status"] == "rejected"


def test_text_edit_refuses_unsaved_slide_text(studio, monkeypatch):
    app, window = studio
    _create_one_slide_project(app, window)
    from presentation_maker_offline import qt_ui
    messages = []
    monkeypatch.setattr(qt_ui.QMessageBox, "information", lambda *args: messages.append(args[-1]))
    window.slide_text.setPlainText("尚未套用的新內容")
    window.edit_instruction.setPlainText("縮短這一頁")
    revision = window.revision
    window.generate_text_edit_candidate()
    assert window._edit_worker is None
    assert window.revision == revision
    assert "先按" in messages[0]


def test_saved_text_mark_can_drive_candidate_and_is_completed_only_on_accept(studio):
    app, window = studio
    _create_one_slide_project(app, window)
    slide = window.document["slides"][0]
    element = next(item for item in slide["elements"] if item["type"] == "text")
    mark = {"id": "mark-1", "slide_id": slide["id"], "element_id": element["id"],
            "rect": [.1, .2, .3, .4], "comment": "改得更簡短", "status": "待處理",
            "base_revision": window.revision}
    window.document["annotations"].append(mark)
    assert window.save_project()
    window._refresh_annotations(slide["id"])
    window.annotation_list.setCurrentRow(0)
    window.use_selected_annotation_for_text_edit()
    assert window.edit_instruction.toPlainText() == "改得更簡短"
    assert window._edit_annotation_id == mark["id"]
    window._receive_text_edit_candidate({
        "slide_id": slide["id"], "element_id": element["id"], "base_revision": window.revision,
        "instruction": "改得更簡短", "original_text": element["text"],
        "proposed_text": "精簡內容", "runtime": "test-local", "annotation_id": mark["id"],
    })
    assert mark["status"] == "待處理"
    window.accept_selected_text_edit()
    saved = window.store.load_document(window.project_id)[1]
    assert saved["slides"][0]["elements"][0]["text"] == "精簡內容"
    assert saved["annotations"][0]["status"] == "已處理"
