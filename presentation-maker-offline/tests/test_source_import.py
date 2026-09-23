from pptx import Presentation
from docx import Document
from pypdf import PdfWriter

from presentation_maker_offline.source_import import import_source


def test_import_text_keeps_page_structure_and_source_path(tmp_path):
    source = tmp_path / "brief.txt"
    source.write_text("封面\n主旨\f第二頁\n補充內容", encoding="utf-8")
    slides, record = import_source(source)
    assert [slide["title"] for slide in slides] == ["封面", "第二頁"]
    assert slides[0]["elements"][0]["text"] == "主旨"
    assert record["path"] == str(source.resolve())
    assert record["role"] == "primary"


def test_import_pptx_preserves_text_and_normalized_geometry(tmp_path):
    source = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度摘要"
    body = slide.placeholders[1]
    body.text = "營收成長\n成本穩定"
    presentation.save(source)

    slides, _record = import_source(source)
    assert slides[0]["title"] == "季度摘要"
    assert slides[0]["elements"][0]["text"] == "營收成長\n成本穩定"
    assert 0 < slides[0]["elements"][0]["x"] < 1
    assert 0 < slides[0]["elements"][0]["width"] < 1


def test_import_docx_extracts_paragraphs(tmp_path):
    source = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("執行摘要")
    doc.add_paragraph("本季營收成長。")
    doc.save(source)
    slides, _record = import_source(source)
    assert slides[0]["title"] == "執行摘要"
    assert slides[0]["elements"][0]["text"] == "本季營收成長。"


def test_scanned_pdf_is_explicitly_flagged_for_ocr(tmp_path):
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open("wb") as output:
        writer.write(output)
    slides, record = import_source(source)
    assert len(slides) == 1
    assert record["warnings"] == ["PDF 第 1 頁沒有可擷取文字；可能需要本機 OCR"]
