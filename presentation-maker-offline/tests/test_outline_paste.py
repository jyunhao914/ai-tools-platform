import pytest
import tkinter as tk
from pathlib import Path
from pptx import Presentation

from presentation_maker_offline.export import export_project_pptx
from presentation_maker_offline.project_document import new_project_document
from presentation_maker_offline.source_import import parse_outline_text
from presentation_maker_offline.workflow import ProjectStore


def test_parse_markdown_outline_into_slides_with_traceable_source():
    title, slides, source = parse_outline_text(
        "# AI 導入策略\n"
        "簡報目的：說明導入步驟\n"
        "## 現況與挑戰\n"
        "- 資料分散\n"
        "- 流程耗時\n"
        "### 補充觀察\n"
        "跨部門資訊不同步\n"
        "## 推動計畫\n"
        "1. 先盤點需求\n"
        "2. 小規模試行\n"
    )

    assert title == "AI 導入策略"
    assert [slide["title"] for slide in slides] == ["現況與挑戰", "推動計畫"]
    assert "簡報目的：說明導入步驟" in slides[0]["elements"][0]["text"]
    assert "• 資料分散" in slides[0]["elements"][0]["text"]
    assert "補充觀察" in slides[0]["elements"][0]["text"]
    assert "• 先盤點需求" in slides[1]["elements"][0]["text"]
    assert source["origin"] == "pasted_text"
    assert len(source["content_sha256"]) == 64


def test_parse_plain_text_outline_by_blank_lines_and_numbered_pages():
    _, plain_slides, _ = parse_outline_text("市場背景\n• 使用者增加\n\n解決方案\n• 建立入口")
    assert [slide["title"] for slide in plain_slides] == ["市場背景", "解決方案"]
    assert plain_slides[1]["elements"][0]["text"] == "• 建立入口"

    _, numbered_slides, _ = parse_outline_text("第 1 頁：市場背景\n需求成長\n第 2 頁：解決方案\n導入流程")
    assert [slide["title"] for slide in numbered_slides] == ["市場背景", "解決方案"]
    assert numbered_slides[0]["elements"][0]["text"] == "需求成長"


def test_parse_outline_with_adjacent_page_marker_stray_slash_and_markdown_table():
    text = (
        "**認識大腸癌｜21頁簡報大綱**\n\n"
        "# 第1頁｜封面：認識大腸癌\n- 副標題：守護腸道健康\n\n"
        "# 第7頁｜瘜肉不等於大腸癌\n- 大腸鏡可切除或取樣第8頁｜無法改變的危險因子\n- 年齡增加\n\n"
        "# 第20頁｜破解常見迷思\n"
        "| **常見迷思正確觀念** |                    |\n"
        "| ------------ | ------------------ |\n"
        "| 沒有症狀就不用檢查 | 早期大腸癌可能沒有症狀 |\n"
        "# 第21頁｜守護腸道健康\n"
        "- \\\n"
    )

    title, slides, _ = parse_outline_text(text)
    assert title == "認識大腸癌｜21頁簡報大綱"
    assert [slide["title"] for slide in slides] == [
        "封面：認識大腸癌", "瘜肉不等於大腸癌", "無法改變的危險因子",
        "破解常見迷思", "守護腸道健康",
    ]
    assert "第8頁" not in slides[1]["elements"][0]["text"]
    assert "正確觀念：早期大腸癌可能沒有症狀" in slides[3]["elements"][0]["text"]
    assert not slides[4]["elements"]

    twenty_one_pages = ["**認識大腸癌｜21頁簡報大綱**"]
    for page in range(1, 22):
        if page == 8:
            continue
        twenty_one_pages.append(f"# 第{page}頁｜第{page}頁標題")
        if page == 7:
            twenty_one_pages.append("- 最後一點第8頁｜第8頁標題\n- 第8頁內容")
        else:
            twenty_one_pages.append("- 頁面內容")
    _, parsed_twenty_one, _ = parse_outline_text("\n\n".join(twenty_one_pages))
    assert len(parsed_twenty_one) == 21
    assert parsed_twenty_one[7]["title"] == "第8頁標題"


def test_pasted_outline_persists_and_exports_as_editable_pptx(tmp_path):
    title, slides, source = parse_outline_text("## 開場\n- 重點一\n## 結論\n- 重點二")
    document = new_project_document(title)
    document["slides"] = slides
    document["sources"] = [source]
    document["outline"] = [{"id": f"outline-{i}", "slide_id": slide["id"], "title": slide["title"]} for i, slide in enumerate(slides)]
    for slide in slides:
        slide["source_id"] = source["id"]
    document["project_id"] = "pasted-outline-project"

    store = ProjectStore(tmp_path / "projects.sqlite3")
    store.create(source["path"], project_id=document["project_id"])
    store.save_document(document["project_id"], document, expected_revision=0)
    loaded = store.load_document(document["project_id"])[1]
    assert loaded["slides"][0]["elements"][0]["text"] == "• 重點一"
    assert loaded["sources"][0]["origin"] == "pasted_text"

    output = export_project_pptx(tmp_path / "outline.pptx", loaded)
    assert output.is_file()
    reopened = Presentation(output)
    assert len(reopened.slides) == 2
    slide_text = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in reopened.slides]
    assert "開場" in slide_text[0] and "重點一" in slide_text[0]
    assert "結論" in slide_text[1] and "重點二" in slide_text[1]


def test_outline_paste_dialog_pastes_previews_and_creates_project(tmp_path, monkeypatch):
    from presentation_maker_offline.ui import PresentationMakerApp

    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    try:
        app = PresentationMakerApp()
    except tk.TclError as exc:
        pytest.skip(f"Tk display is unavailable: {exc}")

    try:
        app.update()
        app.paste_outline()
        app.update()
        dialog = next(widget for widget in app.winfo_children() if widget.winfo_class() == "Toplevel")
        assert dialog.winfo_viewable()
        assert app.grab_current() is None

        def descendants(widget):
            return [widget, *(child for item in widget.winfo_children() for child in descendants(item))]

        widgets = descendants(dialog)
        outline_input = next(widget for widget in widgets if widget.winfo_class() == "Text")
        outline_input.focus_set()
        dialog.clipboard_get = lambda: "# UI 貼上驗收\n## 第一頁\n- 重點甲\n## 第二頁\n- 重點乙"
        buttons = {widget.cget("text"): widget for widget in widgets if widget.winfo_class() == "TButton"}

        buttons["從剪貼簿貼上"].invoke()
        app.update()
        assert "## 第一頁" in outline_input.get("1.0", "end")

        buttons["預覽大綱"].invoke()
        app.update()
        preview = next(widget for widget in widgets if widget.winfo_class() == "Listbox")
        assert preview.get(0, "end") == ("01　第一頁", "02　第二頁")

        from presentation_maker_offline import ui
        output_path = tmp_path / "ui-outline.pptx"
        monkeypatch.setattr(ui.filedialog, "asksaveasfilename", lambda **_kwargs: str(output_path))
        monkeypatch.setattr(ui.messagebox, "showinfo", lambda *_args, **_kwargs: None)
        buttons["建立並匯出 PPTX…"].invoke()
        assert app.document["title"] == "UI 貼上驗收"
        assert len(app.document["slides"]) == 2
        assert len(app.project_store.list_projects()) == 1
        app.update()
        assert app.editor_frame.winfo_viewable()
        assert output_path.is_file()
        assert len(Presentation(output_path).slides) == 2
        notebook = next(widget for widget in descendants(app.editor_frame) if widget.winfo_class() == "TNotebook")
        assert len(notebook.tabs()) == 2
        assert all("候選" not in widget.cget("text") for widget in descendants(app.editor_frame) if widget.winfo_class() == "TButton")
    finally:
        app.destroy()


def test_outline_text_supports_command_v_and_right_click_paste(tmp_path, monkeypatch):
    from presentation_maker_offline.ui import PresentationMakerApp

    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    try:
        app = PresentationMakerApp()
    except tk.TclError as exc:
        pytest.skip(f"Tk display is unavailable: {exc}")

    try:
        app.update()
        app.paste_outline()
        app.update()
        dialog = next(widget for widget in app.winfo_children() if widget.winfo_class() == "Toplevel")

        def descendants(widget):
            return [widget, *(child for item in widget.winfo_children() for child in descendants(item))]

        widgets = descendants(dialog)
        outline_input = next(widget for widget in widgets if widget.winfo_class() == "Text")
        dialog.clipboard_get = lambda: "## 快捷鍵貼上\n- 內容一\n## 第二頁\n- 內容二"
        outline_input.focus_force()
        outline_input.event_generate("<Command-v>")
        app.update()
        assert "## 快捷鍵貼上" in outline_input.get("1.0", "end")
        dialog.after(600, app.quit)
        app.mainloop()
        create_button = next(widget for widget in widgets if widget.winfo_class() == "TButton" and widget.cget("text") == "建立專案")
        assert str(create_button.cget("state")) == "normal"

        outline_input.delete("1.0", "end")
        dialog.clipboard_get = lambda: "## 右鍵貼上\n- 內容三"
        assert outline_input.bind("<Button-3>")
        assert outline_input.bind("<Button-2>")
        assert outline_input.bind("<Command-v>")
    finally:
        app.destroy()


def test_home_is_first_screen_and_recent_project_reopens_in_editor(tmp_path, monkeypatch):
    from presentation_maker_offline.ui import PresentationMakerApp

    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    try:
        app = PresentationMakerApp()
    except tk.TclError as exc:
        pytest.skip(f"Tk display is unavailable: {exc}")

    try:
        app.update()
        assert app.home_frame.winfo_viewable()
        assert not app.editor_frame.winfo_viewable()
        app.title_text.set("最近專案測試")
        app.new_sample()
        app.update()
        assert app.editor_frame.winfo_viewable()

        project_id = app.current_project_id
        app.show_home()
        app.update()
        assert app.recent_list.get(0) == "最近專案測試　·　3 頁"
        app.recent_list.selection_set(0)
        app._open_recent_project()
        app.update()
        assert app.current_project_id == project_id
        assert app.document["title"] == "最近專案測試"
        assert app.editor_frame.winfo_viewable()
    finally:
        app.destroy()


@pytest.mark.parametrize("text", ["", "   \n  "])
def test_parse_rejects_empty_outline(text):
    with pytest.raises(ValueError, match="貼上簡報大綱"):
        parse_outline_text(text)
