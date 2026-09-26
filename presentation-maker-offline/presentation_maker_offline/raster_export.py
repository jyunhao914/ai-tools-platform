from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches
from PySide6.QtWidgets import QApplication

from .project_document import validate_project_document
from .qt_ui import SlidePreview, slide_preview_payload


def export_project_image_pptx(path: str | Path, document: dict, *, asset_root: str | Path) -> Path:
    """Export one full-slide image per page using the same Qt canvas as preview."""
    validate_project_document(document)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    app = QApplication.instance()
    owns_app = app is None
    if owns_app:
        app = QApplication([])
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    canvas = SlidePreview()
    try:
        with TemporaryDirectory(prefix="presentation-image-export-", dir=out.parent) as temporary:
            staging = Path(temporary)
            for index in range(len(document["slides"])):
                canvas.set_slide(**slide_preview_payload(document, index, asset_root, strict_assets=True))
                image = staging / f"slide-{index + 1:04d}.png"
                canvas.save_slide_image(image)
                page = presentation.slides.add_slide(presentation.slide_layouts[6])
                page.shapes.add_picture(str(image), 0, 0, presentation.slide_width, presentation.slide_height)
            staged_pptx = staging / "presentation.pptx"
            presentation.save(staged_pptx)
            staged_pptx.replace(out)
    finally:
        canvas.close()
        if owns_app:
            app.quit()
    return out
