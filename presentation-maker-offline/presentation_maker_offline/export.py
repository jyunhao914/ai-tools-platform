import io
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from .layout_design import LAYOUT_NAMES, fit_body_font, layout_rects
from .project_document import validate_project_document


THEMES = {
    "清爽藍": {"paper": "F8FAFC", "accent": "2459A6", "text": "172B4D", "rule": "CFE1FA", "muted": "64748B"},
    "雜誌編輯": {"paper": "FFF9F2", "accent": "B45309", "text": "44403C", "rule": "F3D7B5", "muted": "78716C"},
    "自然療癒": {"paper": "F2F8F1", "accent": "2F7658", "text": "34483C", "rule": "C8DDCB", "muted": "647568"},
    "高科技": {"paper": "111827", "accent": "38BDF8", "text": "E2E8F0", "rule": "334155", "muted": "94A3B8"},
}


def _color(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def export_demo_pptx(path: str | Path, title: str, outline: list[str], style: str = "清爽藍") -> Path:
    out = Path(path); out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    for index, heading in enumerate([title, *outline]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(245, 248, 252)
        box = slide.shapes.add_textbox(Inches(0.9), Inches(1.1), Inches(11.5), Inches(1.2))
        tf = box.text_frame; tf.text = heading; tf.paragraphs[0].font.size = Pt(34 if index == 0 else 28); tf.paragraphs[0].font.bold = True
        sub = slide.shapes.add_textbox(Inches(0.95), Inches(6.65), Inches(11), Inches(.35))
        sub.text_frame.text = f"{style}  ·  離線簡報製作器  ·  {index + 1}/{len(outline)+1}"
        sub.text_frame.paragraphs[0].font.size = Pt(11)
    prs.save(out); return out


def export_project_pptx(path: str | Path, document: dict, style: str = "清爽藍", *, asset_root: str | Path | None = None) -> Path:
    """Export editable slides using a readable, selected visual theme."""
    validate_project_document(document)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    theme = THEMES.get(style, THEMES["清爽藍"])
    for slide_index, slide_doc in enumerate(document["slides"]):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _color(theme["paper"])
        brand_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.12), Inches(7.5))
        brand_bar.fill.solid()
        brand_bar.fill.fore_color.rgb = _color(theme["accent"])
        brand_bar.line.fill.background()
        is_cover = slide_index == 0
        has_images = any(element.get("type") == "image" for element in slide_doc.get("elements", []))
        layout_name = slide_doc.get("layout")
        title_rect = (layout_rects(layout_name, cover=is_cover, has_image=has_images)["title"]
                      if layout_name in LAYOUT_NAMES else None)
        title_box = slide.shapes.add_textbox(
            Inches(title_rect[0] * 13.333 if title_rect else .95 if is_cover else .8),
            Inches(title_rect[1] * 7.5 if title_rect else 1.35 if is_cover else .55),
            Inches(title_rect[2] * 13.333 if title_rect else 5.5 if is_cover and has_images else 11.4 if is_cover else 11.7),
            Inches(title_rect[3] * 7.5 if title_rect else 1.45 if is_cover else 1.0),
        )
        title_frame = title_box.text_frame
        title_frame.clear()
        title_frame.word_wrap = True
        title_frame.margin_left = 0
        title_frame.margin_right = 0
        title_frame.margin_top = 0
        title_frame.margin_bottom = 0
        title_run = title_frame.paragraphs[0].add_run()
        title_run.text = slide_doc.get("title", "")
        title_run.font.name = "Aptos Display"
        title_run.font.size = Pt(38 if is_cover else (28 if len(title_run.text) < 28 else 24))
        title_run.font.bold = True
        title_run.font.color.rgb = _color(theme["accent"])
        if not is_cover:
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.8), Inches(1.31 if title_rect else 1.62), Inches(11.7), Inches(.035))
            rule.fill.solid()
            rule.fill.fore_color.rgb = _color(theme["rule"])
            rule.line.fill.background()
        for element in slide_doc.get("elements", []):
            x = max(0, min(.98, float(element.get("x", .08))))
            y = max(0, min(.95, float(element.get("y", .24))))
            width = max(.02, min(1-x, float(element.get("width", .84))))
            height = max(.02, min(1-y, float(element.get("height", .58))))
            if element.get("type", "text") == "image":
                relative_path = element.get("asset_path")
                if not relative_path:
                    continue
                if not asset_root:
                    raise ValueError("asset_root is required to export project images")
                root = Path(asset_root).resolve()
                image_path = (root / relative_path).resolve()
                if root not in image_path.parents:
                    raise ValueError("image asset path must stay inside the project asset folder")
                if not image_path.is_file():
                    raise FileNotFoundError(f"找不到專案圖片素材：{relative_path}")
                from PIL import Image
                with Image.open(image_path) as image:
                    image_ratio = image.width / image.height
                box_width, box_height = width * 13.333, height * 7.5
                if element.get("fit") == "crop":
                    target_ratio = box_width / box_height
                    with Image.open(image_path) as image:
                        source_width, source_height = image.size
                        if image_ratio > target_ratio:
                            crop_width = round(source_height * target_ratio)
                            left = (source_width - crop_width) // 2
                            cropped = image.crop((left, 0, left + crop_width, source_height))
                        else:
                            crop_height = round(source_width / target_ratio)
                            top = (source_height - crop_height) // 2
                            cropped = image.crop((0, top, source_width, top + crop_height))
                        image_bytes = io.BytesIO()
                        cropped.convert("RGB").save(image_bytes, format="PNG")
                        image_bytes.seek(0)
                    slide.shapes.add_picture(image_bytes, Inches(x * 13.333), Inches(y * 7.5),
                                             Inches(box_width), Inches(box_height))
                    continue
                if image_ratio > box_width / box_height:
                    draw_width, draw_height = box_width, box_width / image_ratio
                else:
                    draw_height, draw_width = box_height, box_height * image_ratio
                slide.shapes.add_picture(
                    str(image_path), Inches(x * 13.333 + (box_width-draw_width)/2),
                    Inches(y * 7.5 + (box_height-draw_height)/2), Inches(draw_width), Inches(draw_height),
                )
                continue
            if element.get("type", "text") != "text":
                continue
            box = slide.shapes.add_textbox(
                Inches(x * 13.333), Inches(y * 7.5), Inches(width * 13.333), Inches(height * 7.5),
            )
            frame = box.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = Inches(.03)
            frame.margin_right = Inches(.03)
            frame.margin_top = Inches(.02)
            frame.margin_bottom = Inches(.02)
            text = element.get("text", "").strip()
            font_size, _fits = fit_body_font(text, (x, y, width, height), cover=is_cover)
            for line_index, line in enumerate(text.splitlines() or [text]):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line.strip()
                paragraph.space_after = Pt(12 if font_size >= 20 else 8)
                paragraph.line_spacing = 1.12
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = _color(theme["text"])
        footer = slide.shapes.add_textbox(Inches(.8), Inches(7.08), Inches(11.7), Inches(.22))
        footer.text_frame.text = f"{style}　·　{slide_index + 1:02d} / {len(document['slides']):02d}"
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        for run in footer.text_frame.paragraphs[0].runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.font.color.rgb = _color(theme["muted"])
    presentation.save(out)
    return out
