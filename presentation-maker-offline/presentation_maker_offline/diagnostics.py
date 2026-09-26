from __future__ import annotations

import json
import tempfile
from pathlib import Path


def run_self_test(report_path: str, *, image_smoke: bool = False) -> int:
    """Exercise the frozen app without desktop interaction; write a local JSON report."""
    report: dict = {"ok": False, "checks": {}}
    try:
        import torch
        from diffusers import QwenImage21Pipeline  # noqa: F401
        from PIL import Image
        from pptx import Presentation
        from PySide6.QtWidgets import QApplication

        from .export import export_project_pptx
        from .layout_design import apply_slide_layout
        from .qt_ui import PresentationStudio
        from .raster_export import export_project_image_pptx

        report["checks"]["image_runtime_import"] = True
        report["checks"]["torch_mps_available"] = bool(torch.backends.mps.is_available())
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory(prefix="presentation-studio-smoke-") as temporary:
            root = Path(temporary)
            window = PresentationStudio(root)
            window._start_outline()
            window.outline_input.setPlainText("# 測試簡報\n## 第一頁｜封面\n- 原文保留\n## 第二頁｜配圖\n- 圖片與文字")
            if not window.continue_button.isEnabled():
                raise RuntimeError("大綱貼入後無法繼續")
            window._continue_settings()
            window._create_project()
            if len(window.document["slides"]) != 2:
                raise RuntimeError("專案頁數不正確")
            report["checks"]["outline_to_project"] = True
            assets = root / "projects" / window.project_id / "assets"
            assets.mkdir(parents=True, exist_ok=True)
            Image.new("RGB", (160, 100), (45, 141, 90)).save(assets / "smoke.png")
            second = window.document["slides"][1]
            second["elements"].append({"id": "smoke-image", "type": "image", "asset_path": "smoke.png"})
            apply_slide_layout(second, "左圖右文", cover=False)
            if not window.save_project():
                raise RuntimeError("配圖專案未能保存")
            editable = export_project_pptx(root / "editable.pptx", window.document, asset_root=assets)
            raster = export_project_image_pptx(root / "image.pptx", window.document, asset_root=assets)
            report["checks"]["editable_pages"] = len(Presentation(editable).slides)
            raster_deck = Presentation(raster)
            report["checks"]["image_pages"] = len(raster_deck.slides)
            report["checks"]["image_export_is_picture"] = all(
                len(slide.shapes) == 1 and slide.shapes[0].shape_type == 13 for slide in raster_deck.slides
            )
            if report["checks"]["editable_pages"] != 2 or report["checks"]["image_pages"] != 2 or not report["checks"]["image_export_is_picture"]:
                raise RuntimeError("PowerPoint 輸出內容不符合預期")
            window.close()
            app.processEvents()
        if image_smoke:
            from .backends import LocalQwenImageBackend
            from .manifest import CheckpointManifest

            model_path = Path.home() / ".cache/lm-studio/models/Qwen/Qwen-Image-2.1"
            backend = LocalQwenImageBackend(CheckpointManifest(
                "Qwen-Image-2.1 local", str(model_path), "0" * 64,
                "local model directory", "2.1", "see checkpoint license", "diffusers",
            ), num_inference_steps=8, width=512, height=512)
            health = backend.health()
            report["checks"]["image_model_ready"] = health["ok"]
            if not health["ok"]:
                raise RuntimeError("本機圖片模型無法載入：" + "; ".join(health["errors"]))
            result = backend.generate("A friendly hand-painted illustration of a green park and walking path, no text")
            report["checks"]["real_image_path"] = result["image_path"]
            report["checks"]["real_image_device"] = result["device"]
        report["ok"] = True
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"
    Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if report["ok"] else 1
