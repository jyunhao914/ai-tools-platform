"""Reproducible rendering/export/persistence check for the user outline.

This is an acceptance harness, not a substitute for desktop end-to-end QA.
Only explicitly reviewed local visuals may be supplied.
"""
from pathlib import Path
import argparse
import hashlib
import json
import shutil

from .editorial_scene import attach_scene, cover_scene, comparison_scene, render_scene
from .source_import import parse_outline_text
from .project_document import new_project_document
from .workflow import ProjectStore
from .export import export_project_pptx
from .raster_export import export_project_image_pptx


def run(outline: Path, visual: Path, reviewed_sha256: str, output: Path):
    if hashlib.sha256(visual.read_bytes()).hexdigest() != reviewed_sha256:
        raise ValueError('Visual differs from reviewed image')
    output.mkdir(parents=True, exist_ok=False)
    asset = output / 'colon-concept.png'
    shutil.copyfile(visual, asset)
    title, slides, source = parse_outline_text(outline.read_text())
    if len(slides) != 21:
        raise ValueError('Acceptance fixture must have 21 pages')
    doc = new_project_document(title)
    doc['sources'] = [source]
    # Keep the full original project; only reviewed representative pages use scenes.
    doc['slides'] = slides
    cover = slides[0]
    lines = cover['elements'][0]['text'].splitlines()
    scene = cover_scene(cover['title'], lines[0], '\n'.join(lines[1:]),
                        dict(path=str(asset), review=dict(verdict='accepted',
                             image_sha256=reviewed_sha256,
                             scope='Conceptual cover illustration only; not an anatomy reference')))
    # Retain the original heading including its role prefix without shrinking body text.
    scene['texts'][0]['size'] = 82
    attach_scene(cover, scene)
    comparison = slides[19]
    rows = comparison['elements'][0]['text'].splitlines()
    if len(rows) != 11:
        raise ValueError('Expected a heading and five comparison pairs')
    attach_scene(comparison, comparison_scene(comparison['title'], rows[0],
                 list(zip(rows[1::2], rows[2::2]))))
    store = ProjectStore(output / 'project.sqlite3')
    project_id = store.create(str(outline), doc['project_id'])
    store.save_document(project_id, doc, 0)
    store.db.close()
    reopened = ProjectStore(output / 'project.sqlite3')
    revision, saved = reopened.load_document(project_id)
    reopened.db.close()
    assert revision == 1 and len(saved['slides']) == 21
    assert saved['slides'] == doc['slides']
    # Export the two reviewed representative compositions, not unreviewed pages as final.
    saved['slides'] = [saved['slides'][0], saved['slides'][19]]
    for page, slide in zip((1, 20), saved['slides']):
        render_scene(slide['editorial_scene'], output / f'page-{page:02d}.png')
    editable = export_project_pptx(output / 'representative-editable.pptx', saved)
    raster = export_project_image_pptx(output / 'representative-image.pptx', saved, asset_root=output)
    from pptx import Presentation
    native, image = Presentation(editable), Presentation(raster)
    assert len(native.slides) == len(image.slides) == 2
    for page, slide in zip(native.slides, saved['slides']):
        actual = ''.join(shape.text for shape in page.shapes if shape.has_text_frame)
        expected = ''.join(item['text'] for item in slide['editorial_scene']['texts'])
        assert actual == expected
    assert all(len(page.shapes) == 1 and page.shapes[0].image.size == (1920, 1080)
               for page in image.slides)
    report = dict(parsed_pages=21, saved_reopened=True, representative_pages=[1,20],
                  native_text_verified=True, raster_size=[1920,1080],
                  whole_product_accepted=False, project_id=project_id,
                  remaining=['21-page illustration and design', 'native desktop end-to-end verification',
                             'PowerPoint/Keynote visual roundtrip', 'medical content review'])
    (output / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--outline', type=Path, required=True)
    parser.add_argument('--visual', type=Path, required=True)
    parser.add_argument('--reviewed-sha256', required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(run(args.outline, args.visual, args.reviewed_sha256, args.output), ensure_ascii=False))
