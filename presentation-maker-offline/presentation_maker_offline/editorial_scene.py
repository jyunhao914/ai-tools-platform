"""Exact-text editorial canvas prototype, independent of image model spelling."""
from pathlib import Path
import hashlib
import json
from copy import deepcopy
from PySide6.QtCore import QRectF, Qt
from PySide6.QtGui import QColor, QFont, QFontMetricsF, QImage, QPainter
from PySide6.QtWidgets import QApplication

_application = None


def source_fingerprint(slide: dict) -> str:
    payload = dict(title=slide.get('title'), elements=slide.get('elements', []))
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def attach_scene(slide: dict, scene: dict):
    slide['editorial_scene'] = deepcopy(scene)
    slide['editorial_scene']['source_fingerprint'] = source_fingerprint(slide)


def current_scene(slide: dict) -> dict | None:
    scene = slide.get('editorial_scene')
    if scene and scene.get('source_fingerprint') != source_fingerprint(slide):
        raise ValueError('內容已變更，請重新排版後預覽或匯出；不使用過期設計。')
    return scene


def checked_visual(item: dict) -> QImage:
    """Require explicit review of the exact bytes before composing a final page."""
    path = Path(item['path'])
    review = item.get('review', {})
    if review.get('verdict') != 'accepted':
        raise ValueError('Visual asset has not been accepted')
    if hashlib.sha256(path.read_bytes()).hexdigest() != review.get('image_sha256'):
        raise ValueError('Visual asset changed after review')
    image = QImage(str(path))
    if image.isNull():
        raise ValueError('Visual asset is unreadable')
    return image


def scene_image(scene: dict, *, width=1920, height=1080):
    global _application
    _application = QApplication.instance() or QApplication([])
    image = QImage(width, height, QImage.Format.Format_ARGB32)
    image.fill(QColor(scene['background']))
    painter = QPainter(image)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    painter.scale(width / 1920, height / 1080)
    try:
        for item in scene.get('images', []):
            visual = checked_visual(item)
            area = QRectF(*item['rect'])
            if area.width() <= 0 or area.height() <= 0 or not QRectF(0, 0, 1920, 1080).contains(area):
                raise ValueError('Visual outside canvas')
            scale = min(area.width() / visual.width(), area.height() / visual.height())
            target = QRectF(0, 0, visual.width() * scale, visual.height() * scale)
            target.moveCenter(area.center())
            painter.drawImage(target, visual)
        for item in scene['texts']:
            rect = QRectF(*item['rect'])
            if not QRectF(0, 0, 1920, 1080).contains(rect):
                raise ValueError('Text outside canvas')
            font = QFont('PingFang TC')
            font.setPixelSize(item['size'])
            font.setBold(item.get('bold', False))
            flags = Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignTop | Qt.TextFlag.TextWordWrap
            bounds = QFontMetricsF(font).boundingRect(rect, int(flags), item['text'])
            if bounds.height() > rect.height() or bounds.width() > rect.width():
                raise ValueError(f"Text does not fit: {item['text']}")
            painter.setFont(font)
            painter.setPen(QColor(item.get('color', '#132E37')))
            painter.drawText(rect, int(flags), item['text'])
    finally:
        painter.end()
    return image


def render_scene(scene: dict, output: Path, *, width=1920, height=1080):
    image = scene_image(scene, width=width, height=height)
    if output.exists():
        raise FileExistsError(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    if not image.save(str(output)):
        raise OSError('Cannot save scene')
    return output


def export_scene_slide(presentation, scene: dict):
    """Use the same coordinates for native editable text and image objects."""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    # Run the same fit and asset validation before writing editable objects.
    scene_image(scene)
    page = presentation.slides.add_slide(presentation.slide_layouts[6])
    page.background.fill.solid()
    page.background.fill.fore_color.rgb = RGBColor.from_string(scene['background'].lstrip('#'))
    sx, sy = presentation.slide_width / 1920, presentation.slide_height / 1080
    for item in scene.get('images', []):
        visual = checked_visual(item)
        x, y, w, h = item['rect']
        scale = min(w / visual.width(), h / visual.height())
        iw, ih = visual.width() * scale, visual.height() * scale
        page.shapes.add_picture(str(item['path']), int((x + (w-iw)/2)*sx),
                                int((y + (h-ih)/2)*sy), int(iw*sx), int(ih*sy))
    for item in scene['texts']:
        x, y, w, h = item['rect']
        box = page.shapes.add_textbox(int(x*sx), int(y*sy), int(w*sx), int(h*sy))
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.text = item['text']
        for paragraph in frame.paragraphs:
            paragraph.font.name = 'PingFang TC'
            paragraph.font.size = Pt(item['size'] * presentation.slide_height / 12700 / 1080)
            paragraph.font.bold = item.get('bold', False)
            paragraph.font.color.rgb = RGBColor.from_string(item.get('color', '#132E37').lstrip('#'))
    return page


def cover_scene(title: str, subtitle: str, audience: str, visual: dict) -> dict:
    return dict(background='#FFFCF7', images=[dict(visual, rect=[1000, 120, 840, 840])], texts=[
        dict(text=title, rect=[110, 230, 850, 190], size=112, bold=True),
        dict(text=subtitle, rect=[110, 500, 760, 205], size=48, color='#167B7B'),
        dict(text=audience, rect=[110, 810, 790, 100], size=32),
    ])


def comparison_scene(title: str, heading: str, pairs: list[tuple[str, str]]) -> dict:
    """Arrange paired source statements as rows, never flatten into one paragraph."""
    if not 1 <= len(pairs) <= 5:
        raise ValueError('Comparison supports one to five source pairs per page')
    texts = [dict(text=title, rect=[110, 70, 1700, 120], size=72, bold=True),
             dict(text=heading, rect=[110, 210, 1700, 70], size=32, color='#52686B')]
    row_height = 700 / len(pairs)
    for index, (left, right) in enumerate(pairs):
        y = 305 + index * row_height
        texts.extend([
            dict(text=left, rect=[110, y, 650, row_height-20], size=36, color='#785344'),
            dict(text=right, rect=[865, y, 930, row_height-20], size=36, color='#167B7B', bold=True),
        ])
    return dict(background='#FFFCF7', texts=texts)


def eligibility_scene(title: str, lines: list[str]) -> dict:
    """Preserve supplied wording; use scale and whitespace, not repeated cards."""
    if len(lines) != 3:
        raise ValueError('This composition requires exactly three statements')
    import re
    matches = [re.fullmatch(r'(\d+至\d+歲[，,]?)(.*)', line) for line in lines[:2]]
    if not all(matches):
        raise ValueError('Statements must begin with an age range')
    texts = [dict(text=title, rect=[110, 95, 1680, 115], size=66, bold=True)]
    for index, match in enumerate(matches):
        x = 110 + index * 920
        texts.extend([
            dict(text=match[1], rect=[x, 300, 790, 200], size=100, bold=True, color='#167B7B'),
            dict(text=match[2], rect=[x, 510, 760, 250], size=38),
        ])
    texts.append(dict(text=lines[2], rect=[110, 855, 1660, 115], size=58, bold=True))
    return dict(background='#FFF9EF', texts=texts)
