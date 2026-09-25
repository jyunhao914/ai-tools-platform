from __future__ import annotations

import hashlib
import re
from pathlib import Path
from uuid import uuid4


def import_source(path: str | Path, *, asset_dir: str | Path | None = None) -> tuple[list[dict], dict]:
    """Extract editable text and geometry from supported local source files."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"找不到來源檔案：{source}")
    suffix = source.suffix.lower()
    warnings = []
    if suffix == ".pptx":
        slides = _pptx_slides(source, Path(asset_dir) if asset_dir else None)
    elif suffix == ".txt":
        slides = _text_slides(source.read_text(encoding="utf-8-sig"))
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("匯入 DOCX 需要安裝 python-docx") from exc
        doc = Document(str(source))
        slides = _text_slides("\n".join(p.text for p in doc.paragraphs if p.text.strip()))
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("匯入 PDF 需要安裝 pypdf；掃描頁需另行設定本機 OCR") from exc
        reader = PdfReader(str(source))
        slides = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"PDF 第 {index + 1} 頁沒有可擷取文字；可能需要本機 OCR")
            slides.append(_slide_from_text(text, index, page_label=f"PDF 第 {index + 1} 頁"))
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        asset_path = _copy_asset(source, Path(asset_dir)) if asset_dir else None
        element = {"id": str(uuid4()), "type": "image", "x": .08, "y": .2, "width": .84, "height": .64}
        if asset_path:
            element["asset_path"] = asset_path
        slides = [{"id": str(uuid4()), "order": 0, "title": source.stem, "elements": [element], "source_page": 1}]
    else:
        raise ValueError(f"目前不支援此檔案格式：{suffix or '無副檔名'}")
    if not slides:
        raise ValueError("來源檔案沒有可匯入的頁面或文字")
    record = {"id": str(uuid4()), "path": str(source), "role": "primary", "scope": "project", "version": 1, "format": suffix[1:]}
    if warnings:
        record["warnings"] = warnings
    return slides, record


def parse_outline_text(text: str) -> tuple[str, list[dict], dict]:
    """Parse pasted Markdown or plain-text outline into editable slide records."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise ValueError("請先貼上簡報大綱文字。")

    lines = []
    page_marker = re.compile(r"第\s*\d+\s*頁\s*[｜|:：]")
    page_heading = re.compile(r"^\s*#{1,6}\s*第\s*\d+\s*頁\s*[｜|:：]")
    existing_page_heading = next((page_heading.match(line) for line in normalized.splitlines() if page_heading.match(line)), None)
    page_level = len(re.match(r"^\s*(#+)", existing_page_heading.group(0)).group(1)) if existing_page_heading else 2
    for raw_line in normalized.splitlines():
        marker = page_marker.search(raw_line)
        if marker and marker.start() > 0 and not page_heading.match(raw_line):
            before = raw_line[:marker.start()].rstrip()
            after = raw_line[marker.start():].strip()
            if before:
                lines.append(before)
            # A page marker embedded after body text keeps the document's slide-heading level.
            # Otherwise it may become the deck title or be dropped from mixed-level headings.
            lines.append(f"{'#' * page_level} {after}")
        else:
            lines.append(raw_line)
    heading_rows = []
    for index, line in enumerate(lines):
        match = re.match(r"^\s*(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if match:
            heading_rows.append((index, len(match.group(1)), match.group(2).strip()))

    deck_title = ""
    heading_preface: list[str] = []
    slide_sections: list[tuple[str, list[str]]] = []
    if heading_rows:
        has_deck_heading = any(level == 1 for _, level, _ in heading_rows) and any(level >= 2 for _, level, _ in heading_rows)
        if has_deck_heading:
            deck_row = next((row for row in heading_rows if row[1] == 1), None)
            deck_title = deck_row[2] if deck_row else ""
            slide_rows = [(i, level, title) for i, level, title in heading_rows if level == 2]
            if not slide_rows:
                slide_rows = [(i, level, title) for i, level, title in heading_rows if level >= 2]
            if deck_row and slide_rows:
                heading_preface = lines[deck_row[0] + 1:slide_rows[0][0]]
        else:
            slide_rows = heading_rows
            preface_title = re.match(r"^\*\*(.+?)\*\*\s*$", lines[0].strip()) if lines else None
            deck_title = preface_title.group(1).strip() if preface_title else heading_rows[0][2]
        for row_index, (line_index, _level, title) in enumerate(slide_rows):
            end = slide_rows[row_index + 1][0] if row_index + 1 < len(slide_rows) else len(lines)
            body_lines = list(heading_preface) if row_index == 0 else []
            for body_line in lines[line_index + 1:end]:
                child_heading = re.match(r"^\s*#{1,6}\s+(.+?)\s*#*\s*$", body_line)
                body_lines.append(child_heading.group(1) if child_heading else body_line)
            slide_sections.append((title, body_lines))
    else:
        # Numbered page/slide headings are common in outlines pasted from notes.
        numbered = re.compile(r"^\s*(?:第\s*\d+\s*[頁章節]\s*[、.．:：-]?\s*|\d+\s*[.、．:：]\s*)(.+?)\s*$")
        numbered_rows = [(i, match.group(1)) for i, line in enumerate(lines) if (match := numbered.match(line))]
        if len(numbered_rows) >= 2:
            for row_index, (line_index, title) in enumerate(numbered_rows):
                end = numbered_rows[row_index + 1][0] if row_index + 1 < len(numbered_rows) else len(lines)
                slide_sections.append((title, lines[line_index + 1:end]))
            deck_title = numbered_rows[0][1]
        else:
            # Blank-line-separated blocks become slides; a block's first line is its title.
            blocks = [block.strip() for block in re.split(r"\n\s*\n|\n\s*---+\s*\n|\f", normalized) if block.strip()]
            for block in blocks:
                block_lines = [line.strip() for line in block.splitlines() if line.strip()]
                if block_lines:
                    slide_sections.append((block_lines[0], block_lines[1:]))
            if len(slide_sections) == 1 and len([line for line in lines if line.strip()]) > 5:
                # Retain the existing TXT importer behavior for long unstructured text.
                slide_sections = []
                paragraphs = [line.strip() for line in lines if line.strip()]
                for offset in range(0, len(paragraphs), 5):
                    slide_sections.append((paragraphs[offset], paragraphs[offset + 1:offset + 5]))
            if slide_sections:
                deck_title = slide_sections[0][0]

    slide_sections = [(title.strip(), body) for title, body in slide_sections if title.strip()]
    if not slide_sections:
        raise ValueError("無法從文字辨識投影片；請用 Markdown 標題（## 標題）或空行分隔各頁。")
    if len(slide_sections) > 100:
        raise ValueError("一次最多貼入 100 張投影片的大綱。")

    slides = []
    for index, (title, body_lines) in enumerate(slide_sections):
        cleaned_body = []
        for line in body_lines:
            line = line.strip()
            if re.fullmatch(r"[-*+]\s*\\", line):
                continue
            if line.startswith("|") and re.fullmatch(r"\|?[\s:|\-]+\|?", line):
                continue
            if line.startswith("|") and line.endswith("|"):
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                cells = [re.sub(r"\*\*(.+?)\*\*", r"\1", cell).strip() for cell in cells]
                cells = [cell for cell in cells if cell]
                if len(cells) >= 2:
                    cleaned_body.extend((f"迷思：{cells[0]}", f"正確觀念：{cells[1]}"))
                    continue
                if cells:
                    cleaned_body.append(cells[0])
                    continue
            line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            line = re.sub(r"^\s*(?:[-*+]\s+|\d+[.)、．]\s*)", "• ", line)
            if line:
                cleaned_body.append(line)
        clean_title = re.sub(r"^第\s*\d+\s*頁\s*[｜|:：]\s*", "", title).strip()
        slides.append(_slide_from_text("\n".join([clean_title, *cleaned_body]), index))

    record_id = str(uuid4())
    record = {
        "id": record_id,
        "path": f"paste://{record_id}",
        "display_name": "貼上的簡報大綱",
        "origin": "pasted_text",
        "text": normalized,
        "content_sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "role": "primary",
        "scope": "project",
        "version": 1,
        "format": "text",
    }
    return deck_title or slides[0]["title"], slides, record


def _text_slides(text: str) -> list[dict]:
    sections = [section.strip() for section in text.split("\f") if section.strip()]
    if len(sections) == 1:
        paragraphs = [line.strip() for line in sections[0].splitlines() if line.strip()]
        sections = ["\n".join(paragraphs[i:i + 5]) for i in range(0, len(paragraphs), 5)]
    return [_slide_from_text(section, i) for i, section in enumerate(sections)]


def _slide_from_text(text: str, index: int, page_label: str | None = None) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = lines[0][:120] if lines else (page_label or f"第 {index + 1} 頁")
    body = "\n".join(lines[1:])
    elements = []
    if body:
        elements.append({"id": str(uuid4()), "type": "text", "text": body, "x": 0.08, "y": 0.24, "width": 0.84, "height": 0.58})
    return {"id": str(uuid4()), "order": index, "title": title, "elements": elements, "source_page": index + 1}


def _copy_asset(source: Path, asset_dir: Path) -> str:
    return _copy_asset_bytes(source.read_bytes(), source.suffix, asset_dir)


def _copy_asset_bytes(content: bytes, suffix: str, asset_dir: Path) -> str:
    asset_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{uuid4().hex}{suffix.lower()}"
    (asset_dir / filename).write_bytes(content)
    return filename


def _pptx_slides(path: Path, asset_dir: Path | None = None) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    slides = []
    width, height = presentation.slide_width, presentation.slide_height
    for index, source_slide in enumerate(presentation.slides):
        title_shape = source_slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.text.strip() else f"第 {index + 1} 頁"
        title_shape_id = title_shape.shape_id if title_shape else None
        elements = []
        for shape in source_slide.shapes:
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                element = {
                    "id": str(uuid4()), "type": "image",
                    "x": shape.left / width, "y": shape.top / height,
                    "width": shape.width / width, "height": shape.height / height,
                    "source_shape": shape.name,
                }
                if asset_dir:
                    element["asset_path"] = _copy_asset_bytes(image.blob, f".{image.ext}", asset_dir)
                elements.append(element)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text.strip()
            if not text or shape.shape_id == title_shape_id:
                continue
            elements.append({
                "id": str(uuid4()), "type": "text", "text": text,
                "x": shape.left / width, "y": shape.top / height,
                "width": shape.width / width, "height": shape.height / height,
                "source_shape": shape.name,
            })
        slides.append({"id": str(uuid4()), "order": index, "title": title, "elements": elements, "source_page": index + 1})
    return slides
